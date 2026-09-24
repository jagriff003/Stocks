"""
Tests for the rebalance workbook's structure and its Target fill.

The formulas themselves are verified by recalculating a filled workbook in
Excel and comparing every Plan row with an independent Python implementation
(done when the workbook was built, 2026-09-24: 0 formula errors, 33/33 rows
exact).  Excel is not available to pytest, so these cover what can go wrong
without it: the layout, refusing to overwrite account data, and the fill rules —
above all that a Track K position you chose to hold is never dropped silently.

Run:  python -m pytest tests/test_rebalance_book.py -v
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT))

from momentum.rebalance_book import (BLOCK, P0, P1, T0, T1, build_workbook,  # noqa: E402
                                     fill_target)


def rows():
    return [{"symbol": "MU", "model": "J", "weight": 0.0625, "price": 120.5},
            {"symbol": "gOOgl", "model": "J", "weight": 0.0625, "price": 330.0},
            {"symbol": "PDBC", "model": "K", "weight": 0.25, "price": 14.1}]


def target(path):
    ws = load_workbook(path)["Target"]
    return {ws[f"A{r}"].value: {c: ws[f"{c}{r}"].value for c in "BCDEFG"}
            for r in range(T0, T1 + 1) if ws[f"A{r}"].value}


def test_layout_and_capacity(tmp_path):
    p = build_workbook(tmp_path / "R.xlsx")
    wb = load_workbook(p)
    assert wb.sheetnames == ["Trades", "Setup", "Target", "Holdings", "Plan", "Guide"]
    # the phone list must be able to show every target name, and the Plan every
    # target name plus every held one
    assert BLOCK >= T1 - T0 + 1
    assert P1 - P0 + 1 >= 2 * (T1 - T0 + 1)
    assert wb.calculation.fullCalcOnLoad


def test_refuses_to_overwrite_account_data(tmp_path):
    p = build_workbook(tmp_path / "R.xlsx")
    with pytest.raises(FileExistsError):
        build_workbook(p)
    build_workbook(p, overwrite=True)


def test_fill_defaults_track_k_to_not_taken_and_upper_cases(tmp_path):
    p = build_workbook(tmp_path / "R.xlsx")
    fill_target(rows(), "2026-09-29", "test", p)
    t = target(p)
    assert set(t) == {"MU", "GOOGL", "PDBC"}
    assert t["PDBC"]["E"] == "N" and t["MU"]["E"] is None
    assert all(v["G"] == "Auto" for v in t.values())


def test_fill_keeps_hand_choices_and_never_drops_a_taken_k_row(tmp_path):
    p = build_workbook(tmp_path / "R.xlsx")
    fill_target(rows(), "2026-09-29", "test", p)
    wb = load_workbook(p)
    ws = wb["Target"]
    for r in range(T0, T1 + 1):
        if ws[f"A{r}"].value == "PDBC":
            ws[f"E{r}"], ws[f"F{r}"], ws[f"G{r}"] = "Y", 0.15, "Trad"
        if ws[f"A{r}"].value == "MU":
            ws[f"G{r}"] = "Roth"
    wb.save(p)
    # next rotation: Track K no longer recommends PDBC, MU drops out of Track J
    res = fill_target([{"symbol": "GOOGL", "model": "J", "weight": 1.0, "price": 331.0}],
                      "2026-10-13", "test", p)
    t = target(p)
    assert res["carried_k"] == 1
    assert "PDBC" in t and t["PDBC"]["E"] == "Y" and t["PDBC"]["F"] == 0.15 and t["PDBC"]["G"] == "Trad"
    assert t["PDBC"]["C"] == 0.0                     # the model's weight is gone; yours is kept
    assert "MU" not in t                             # a Track J name that left is sold, as designed
